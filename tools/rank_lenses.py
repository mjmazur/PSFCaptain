"""Ranks lenses from best to worst using the lens_quality_report_metrics.csv files from star_measure.

Usage:
    python tools/rank_lenses.py /path/to/LensTests/25mm_LensTesting

Each subdirectory of the given directory is one lens and must hold the
lens_quality_report_metrics.csv written next to its lens_quality_report.pdf. The per-frame metrics
are reduced to one value per lens (median over frames), each scored 0-100 relative to the other
lenses (best lens = 100, worst = 0), and combined with the weights in METRICS.
"""
import os
import sys
import argparse
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

METRICS_FILE = 'lens_quality_report_metrics.csv'

# column, label, unit, higher is better, weight, description
METRICS = [
    ('fwhm|Center', 'FWHM center', 'px', False, 15, 'Star size at the optical axis: best-case sharpness.'),
    ('fwhm|Field', 'FWHM field', 'px', False, 15, 'Median star size over the whole frame: overall sharpness.'),
    ('ee80|Field', 'EE80 field', 'px', False, 10, 'Radius containing 80% of the star flux; sensitive to halos and wings.'),
    ('fwhm|Corners', 'FWHM corners', 'px', False, 15, 'Mean star size in the four corners: off-axis sharpness.'),
    ('corner_fwhm_spread', 'Corner FWHM spread', '%', False, 10, 'Best-to-worst corner FWHM difference relative to center: tilt/decentering.'),
    ('elongation|Field', 'Elongation field', '', False, 10, 'Star axis ratio (1 = round): coma and astigmatism.'),
    ('mlim|Field', 'Limiting mag field', 'mag', True, 15, '5-sigma limiting magnitude: throughput, PSF size and noise combined.'),
    ('rel_illum|Corners', 'Corner illumination', '%', True, 10, 'Corner flux relative to center: vignetting.'),
]
EXTRA_COLUMNS = [('fwhm_arc|Field', 'FWHM field (arcsec)'), ('stars|Field', 'Stars per frame'),
                 ('radial_distortion', 'Radial distortion k'), ('distortion', 'Scale change corners (%)')]


def find_lens_dirs(root):
    """Maps lens name to metrics CSV path for every subdirectory of root that has one."""
    found, missing = {}, []
    for name in sorted(os.listdir(root)):
        path = os.path.join(root, name)
        if not os.path.isdir(path):
            continue
        csv = os.path.join(path, METRICS_FILE)
        if os.path.exists(csv):
            found[name] = csv
        else:
            missing.append(name)
    if not found and os.path.exists(os.path.join(root, METRICS_FILE)):
        found[os.path.basename(os.path.abspath(root))] = os.path.join(root, METRICS_FILE)
    return found, missing


def summarize(csvs):
    """One row per lens: median over frames of every metric, plus frame count and scatter."""
    rows = {}
    for name, csv in csvs.items():
        df = pd.read_csv(csv).drop(columns='frame', errors='ignore')
        row = df.median()
        row['frames'] = len(df)
        for col, *_ in METRICS:
            if col in df.columns:
                row[f'{col} (std)'] = df[col].std()
        rows[name] = row
    return pd.DataFrame(rows).T


def score(summary, weights):
    """Per-metric 0-100 scores (min-max across lenses) and the weighted total."""
    scores = pd.DataFrame(index=summary.index)
    for col, _, _, higher, _, _ in METRICS:
        if col not in summary.columns or summary[col].notna().sum() == 0:
            continue
        v = summary[col].astype(float)
        span = v.max() - v.min()
        s = (v - v.min()) / span if span > 0 else pd.Series(1.0, index=v.index)
        scores[col] = 100 * (s if higher else 1 - s)
    w = pd.Series({c: weights[c] for c in scores.columns})
    # Weighted mean over the metrics a lens actually has, so a missing value does not count as zero.
    scores['total'] = (scores[w.index] * w).sum(axis=1, skipna=True) / (scores[w.index].notna() * w).sum(axis=1)
    return scores


def build_tables(summary, scores):
    # Ties on the total are broken by the sharpest center FWHM.
    order = pd.DataFrame({'total': scores['total'].round(1), 'fwhm': summary['fwhm|Center'].astype(float)})
    ranked = order.sort_values(['total', 'fwhm'], ascending=[False, True]).index
    ranking = pd.DataFrame(index=ranked)
    ranking.insert(0, 'Rank', range(1, len(ranked) + 1))
    ranking['Score'] = scores.loc[ranked, 'total'].round(1)
    ranking['Frames'] = summary.loc[ranked, 'frames'].astype(int)
    for col, label, unit, *_ in METRICS:
        if col in summary.columns:
            ranking[f'{label} ({unit})' if unit else label] = summary.loc[ranked, col]
    for col, label in EXTRA_COLUMNS:
        if col in summary.columns:
            ranking[label] = summary.loc[ranked, col]
    for col, label, *_ in METRICS:
        if col in scores.columns:
            ranking[f'Score: {label}'] = scores.loc[ranked, col].round(1)
    ranking.index.name = 'Lens'

    metric_cols = [c for c in summary.columns if not c.endswith('(std)') and c != 'frames']
    return ranking.reset_index(), summary.loc[ranked, metric_cols].rename_axis('Lens').reset_index()


def method_table(weights):
    rows = [{'Metric': label, 'Column': col, 'Unit': unit, 'Better': 'higher' if higher else 'lower',
             'Weight (%)': weights[col], 'Meaning': desc}
            for col, label, unit, higher, _, desc in METRICS]
    return pd.DataFrame(rows)


def style_sheet(ws, header_fill='1F3864'):
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=header_fill)
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = 'C2'
    for i, col in enumerate(ws.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[1:] or col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 14 if i > 1 else 20)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = '0' if c.value.is_integer() else '0.000'


def write_workbook(path, ranking, metrics, method, root):
    with pd.ExcelWriter(path, engine='openpyxl') as xw:
        ranking.to_excel(xw, sheet_name='Ranking', index=False)
        metrics.to_excel(xw, sheet_name='All metrics', index=False)
        method.to_excel(xw, sheet_name='Method', index=False)
        for name in ('Ranking', 'All metrics'):
            style_sheet(xw.sheets[name])
        ws = xw.sheets['Ranking']
        for row in range(2, ws.max_row + 1):
            ws.cell(row, 3).font = Font(bold=True)
            ws.cell(row, 3).number_format = '0.0'
            if row == 2:
                for c in ws[row][:3]:
                    c.fill = PatternFill('solid', fgColor='C6E0B4')
        ws = xw.sheets['Method']
        style_sheet(ws, '595959')
        ws.freeze_panes = 'A2'
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 24
        ws.column_dimensions['F'].width = 80
        n = len(method) + 3
        notes = [f'Source: {root}',
                 'Each metric is the median over all frames of a lens, then scaled 0-100 across the lenses '
                 'ranked here (best = 100, worst = 0).',
                 'Score = weighted mean of the metric scores. Scores are relative: adding or removing a lens '
                 'changes them.']
        for i, text in enumerate(notes):
            ws.cell(n + i, 1, text)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('directory', help=f'Directory with one subdirectory per lens, each holding {METRICS_FILE}.')
    parser.add_argument('-o', '--output', help='Output .xlsx (default: <directory>/lens_ranking.xlsx).')
    args = parser.parse_args()

    root = os.path.abspath(args.directory)
    if not os.path.isdir(root):
        sys.exit(f'Not a directory: {root}')
    csvs, missing = find_lens_dirs(root)
    if not csvs:
        sys.exit(f'No subdirectories of {root} contain {METRICS_FILE}; run star_measure first.')
    for name in missing:
        print(f'Skipping {name}: no {METRICS_FILE}')

    weights = {col: w for col, _, _, _, w, _ in METRICS}
    summary = summarize(csvs)
    scores = score(summary, weights)
    ranking, metrics = build_tables(summary, scores)

    output = args.output or os.path.join(root, 'lens_ranking.xlsx')
    write_workbook(output, ranking, metrics, method_table(weights), root)

    print(f'{"Rank":<5}{"Lens":<24}{"Score":>6}{"FWHM ctr":>10}{"FWHM corners":>14}{"mlim":>7}')
    for _, r in ranking.iterrows():
        print(f'{r["Rank"]:<5}{r["Lens"]:<24}{r["Score"]:>6.1f}{r["FWHM center (px)"]:>10.2f}'
              f'{r["FWHM corners (px)"]:>14.2f}{r["Limiting mag field (mag)"]:>7.2f}')
    print(f'Wrote {output}')


if __name__ == '__main__':
    main()
